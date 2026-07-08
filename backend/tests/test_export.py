import io

import openpyxl

from app.export import series_to_csv, series_to_excel


def test_series_to_csv_includes_series_column_and_all_points():
    series = [
        {"name": "Arm A", "points": [(0.0, 1.0), (2.0, 0.95)]},
        {"name": "Arm B", "points": [(0.0, 1.0), (2.0, 0.90)]},
    ]

    csv_text = series_to_csv(series)

    lines = csv_text.strip().splitlines()
    assert lines[0] == "series,x,y"
    assert "Arm A,0.0,1.0" in csv_text
    assert "Arm B,2.0,0.9" in csv_text
    assert len(lines) == 5  # header + 4 data rows


def test_series_to_excel_produces_readable_workbook():
    series = [{"name": "Series A", "points": [(1.0, 13.0), (5.0, 45.0)]}]

    excel_bytes = series_to_excel(series)

    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("series", "x", "y")
    assert rows[1] == ("Series A", 1.0, 13.0)
    assert rows[2] == ("Series A", 5.0, 45.0)


def test_series_to_csv_escapes_formula_injection_in_series_name():
    series = [{"name": "=1+1", "points": [(0.0, 1.0)]}]

    csv_text = series_to_csv(series)

    lines = csv_text.strip().splitlines()
    assert lines[1] == "'=1+1,0.0,1.0"
    assert "=1+1,0.0,1.0" not in csv_text.replace("'=1+1,0.0,1.0", "")


def test_series_to_excel_escapes_formula_injection_in_series_name():
    series = [{"name": "=cmd|'/c calc'!A1", "points": [(0.0, 1.0)]}]

    excel_bytes = series_to_excel(series)

    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][0] == "'=cmd|'/c calc'!A1"
    assert not rows[1][0].startswith("=")
